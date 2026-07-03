#!/usr/bin/env python3
"""PDF Dönüştürücü - PDF dosyalarını Word (.docx) veya Excel (.xlsx) formatına çevirir.

Grafik arayüz (Tkinter) ile kullanım:
    python3 pdf_donusturucu.py

Komut satırından kullanım:
    python3 pdf_donusturucu.py dosya.pdf --format word
    python3 pdf_donusturucu.py dosya.pdf --format excel --cikti /hedef/klasor
"""

import argparse
import os
import sys
import threading
import traceback
import unicodedata
from pathlib import Path


def _metin_duzelt(metin):
    """Çıkarılan metni temizler ve Türkçe karakterleri düzgün hale getirir.

    Bazı PDF'ler 'ğ', 'ü', 'ş' gibi harfleri ayrık aksan işaretleriyle kodlar;
    NFC normalizasyonu bunları tek karaktere birleştirir. Ayrıca Word/Excel'in
    kabul etmediği kontrol karakterleri ayıklanır.
    """
    if not metin:
        return ""
    metin = unicodedata.normalize("NFC", metin)
    return "".join(k for k in metin
                   if k in "\n\t" or unicodedata.category(k)[0] != "C")


# ---------------------------------------------------------------------------
# OCR (taranmış PDF'ler için)
# ---------------------------------------------------------------------------

_OCR_DILI = None  # ilk kullanımda belirlenir; False = OCR kullanılamıyor


def _tesseract_hazirla():
    """Tesseract'ı bulur ve kullanılacak dili döndürür (yoksa False).

    PyInstaller paketinde çalışıyorsak paket içindeki tesseract kullanılır,
    aksi halde sistemdeki aranır. Türkçe dil paketi varsa Türkçe+İngilizce,
    yoksa yalnızca İngilizce seçilir.
    """
    global _OCR_DILI
    if _OCR_DILI is not None:
        return _OCR_DILI

    import shutil

    try:
        import pytesseract
    except ImportError:
        _OCR_DILI = False
        return False

    paket_dizini = getattr(sys, "_MEIPASS", None)
    if paket_dizini:
        aday = Path(paket_dizini) / "tesseract" / "tesseract"
        if aday.exists():
            pytesseract.pytesseract.tesseract_cmd = str(aday)
            os.environ["TESSDATA_PREFIX"] = str(
                Path(paket_dizini) / "tesseract" / "tessdata")

    if (not paket_dizini or not (Path(paket_dizini) / "tesseract" /
                                 "tesseract").exists()) \
            and not shutil.which("tesseract"):
        _OCR_DILI = False
        return False

    try:
        diller = set(pytesseract.get_languages(config=""))
    except Exception:
        _OCR_DILI = False
        return False

    if "tur" in diller:
        _OCR_DILI = "tur+eng" if "eng" in diller else "tur"
    elif "eng" in diller:
        _OCR_DILI = "eng"
    else:
        _OCR_DILI = False
    return _OCR_DILI


def _sayfa_ocr(sayfa):
    """Sayfayı görüntüye çevirip OCR ile metnini okur. OCR yoksa '' döner."""
    dil = _tesseract_hazirla()
    if not dil:
        return ""
    import pytesseract
    goruntu = sayfa.to_image(resolution=300).original
    return _metin_duzelt(pytesseract.image_to_string(goruntu, lang=dil))


def ocr_kullanilabilir():
    """OCR motorunun kullanılabilir olup olmadığını döndürür."""
    return bool(_tesseract_hazirla())


# ---------------------------------------------------------------------------
# Dönüştürme mantığı
# ---------------------------------------------------------------------------

def _tablo_bolgeleri(page):
    """Sayfadaki tabloların kapladığı dikdörtgen bölgeleri döndürür."""
    bolgeler = []
    for tablo in page.find_tables():
        bolgeler.append(tablo.bbox)
    return bolgeler


def _tablo_disindaki_metin(page, bolgeler):
    """Tablo bölgelerinin dışında kalan sayfa metnini döndürür."""
    if not bolgeler:
        return _metin_duzelt(page.extract_text() or "")

    def tablo_disinda(obj):
        merkez_y = (obj["top"] + obj["bottom"]) / 2
        merkez_x = (obj["x0"] + obj["x1"]) / 2
        for (x0, top, x1, bottom) in bolgeler:
            if x0 <= merkez_x <= x1 and top <= merkez_y <= bottom:
                return False
        return True

    filtreli = page.filter(tablo_disinda)
    return _metin_duzelt(filtreli.extract_text() or "")


def pdf_to_word(pdf_yolu, docx_yolu, ilerleme=None, ocr=True):
    """PDF dosyasını Word belgesine çevirir. Metinler paragraf, tablolar tablo olarak aktarılır."""
    import pdfplumber
    from docx import Document
    from docx.shared import Pt

    belge = Document()

    with pdfplumber.open(pdf_yolu) as pdf:
        toplam = len(pdf.pages)
        for sayfa_no, sayfa in enumerate(pdf.pages, start=1):
            if ilerleme:
                ilerleme(sayfa_no, toplam)

            bolgeler = _tablo_bolgeleri(sayfa)
            metin = _tablo_disindaki_metin(sayfa, bolgeler)

            # Sayfada okunabilir metin yoksa taranmış olabilir → OCR dene
            if ocr and not bolgeler and len(metin.strip()) < 20:
                ocr_metni = _sayfa_ocr(sayfa)
                if len(ocr_metni.strip()) > len(metin.strip()):
                    metin = ocr_metni

            if metin.strip():
                for satir in metin.splitlines():
                    if satir.strip():
                        belge.add_paragraph(satir)

            for tablo in sayfa.extract_tables():
                if not tablo:
                    continue
                satir_sayisi = len(tablo)
                sutun_sayisi = max(len(satir) for satir in tablo)
                w_tablo = belge.add_table(rows=satir_sayisi, cols=sutun_sayisi)
                w_tablo.style = "Table Grid"
                for i, satir in enumerate(tablo):
                    for j, hucre in enumerate(satir):
                        w_tablo.cell(i, j).text = _metin_duzelt(
                            (hucre or "").strip())
                belge.add_paragraph()

            if sayfa_no < toplam:
                belge.add_page_break()

    belge.save(docx_yolu)


def pdf_to_excel(pdf_yolu, xlsx_yolu, ilerleme=None, ocr=True):
    """PDF dosyasını Excel çalışma kitabına çevirir.

    Her sayfa ayrı bir çalışma sayfasına yazılır. Tablolar hücrelere dağıtılır,
    tablo dışındaki metinler satır satır eklenir.
    """
    import pdfplumber
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    kitap = Workbook()
    kitap.remove(kitap.active)

    with pdfplumber.open(pdf_yolu) as pdf:
        toplam = len(pdf.pages)
        for sayfa_no, sayfa in enumerate(pdf.pages, start=1):
            if ilerleme:
                ilerleme(sayfa_no, toplam)

            ws = kitap.create_sheet(title=f"Sayfa {sayfa_no}")
            aktif_satir = 1

            bolgeler = _tablo_bolgeleri(sayfa)
            metin = _tablo_disindaki_metin(sayfa, bolgeler)

            # Sayfada okunabilir metin yoksa taranmış olabilir → OCR dene
            if ocr and not bolgeler and len(metin.strip()) < 20:
                ocr_metni = _sayfa_ocr(sayfa)
                if len(ocr_metni.strip()) > len(metin.strip()):
                    metin = ocr_metni

            if metin.strip():
                for satir in metin.splitlines():
                    if satir.strip():
                        ws.cell(row=aktif_satir, column=1, value=satir.strip())
                        aktif_satir += 1
                aktif_satir += 1

            for tablo in sayfa.extract_tables():
                if not tablo:
                    continue
                baslik_satiri = aktif_satir
                for satir in tablo:
                    for j, hucre in enumerate(satir, start=1):
                        ws.cell(row=aktif_satir, column=j,
                                value=_metin_duzelt((hucre or "").strip()))
                    aktif_satir += 1
                for hucre in ws[baslik_satiri]:
                    hucre.font = Font(bold=True)
                aktif_satir += 1

            # Sütun genişliklerini içeriğe göre ayarla
            for sutun in ws.columns:
                en_uzun = max((len(str(h.value)) for h in sutun if h.value),
                              default=0)
                harf = get_column_letter(sutun[0].column)
                ws.column_dimensions[harf].width = min(max(en_uzun + 2, 10), 80)

    if not kitap.sheetnames:
        kitap.create_sheet(title="Sayfa 1")
    kitap.save(xlsx_yolu)


def donustur(pdf_yolu, format_secimi, cikti_klasoru=None, ilerleme=None,
             ocr=True):
    """Tek bir PDF dosyasını seçilen formata çevirir, oluşan dosyanın yolunu döndürür."""
    pdf_yolu = Path(pdf_yolu)
    if not pdf_yolu.exists():
        raise FileNotFoundError(f"Dosya bulunamadı: {pdf_yolu}")

    klasor = Path(cikti_klasoru) if cikti_klasoru else pdf_yolu.parent
    klasor.mkdir(parents=True, exist_ok=True)

    if format_secimi == "word":
        hedef = klasor / (pdf_yolu.stem + ".docx")
        pdf_to_word(str(pdf_yolu), str(hedef), ilerleme, ocr)
    elif format_secimi == "excel":
        hedef = klasor / (pdf_yolu.stem + ".xlsx")
        pdf_to_excel(str(pdf_yolu), str(hedef), ilerleme, ocr)
    else:
        raise ValueError(f"Bilinmeyen format: {format_secimi}")

    return hedef


# ---------------------------------------------------------------------------
# Grafik arayüz
# ---------------------------------------------------------------------------

def arayuz_baslat():
    import tkinter as tk
    from tkinter import filedialog, messagebox, ttk

    pencere = tk.Tk()
    pencere.title("PDF Dönüştürücü")
    pencere.geometry("620x480")
    pencere.minsize(520, 420)

    secili_dosyalar = []

    ana = ttk.Frame(pencere, padding=15)
    ana.pack(fill="both", expand=True)

    baslik = ttk.Label(ana, text="PDF → Word / Excel Dönüştürücü",
                       font=("", 14, "bold"))
    baslik.pack(pady=(0, 10))

    # --- Dosya listesi ---
    liste_cerceve = ttk.LabelFrame(ana, text="Seçilen PDF dosyaları", padding=8)
    liste_cerceve.pack(fill="both", expand=True)

    liste = tk.Listbox(liste_cerceve, selectmode="extended")
    liste.pack(side="left", fill="both", expand=True)
    kaydirma = ttk.Scrollbar(liste_cerceve, command=liste.yview)
    kaydirma.pack(side="right", fill="y")
    liste.config(yscrollcommand=kaydirma.set)

    def dosya_sec():
        dosyalar = filedialog.askopenfilenames(
            title="PDF dosyalarını seçin",
            filetypes=[("PDF dosyaları", "*.pdf"), ("Tüm dosyalar", "*.*")])
        for d in dosyalar:
            if d not in secili_dosyalar:
                secili_dosyalar.append(d)
                liste.insert("end", os.path.basename(d))

    def secilenleri_kaldir():
        for indeks in reversed(liste.curselection()):
            liste.delete(indeks)
            del secili_dosyalar[indeks]

    dugme_satiri = ttk.Frame(ana)
    dugme_satiri.pack(fill="x", pady=6)
    ttk.Button(dugme_satiri, text="➕ PDF Ekle",
               command=dosya_sec).pack(side="left")
    ttk.Button(dugme_satiri, text="➖ Kaldır",
               command=secilenleri_kaldir).pack(side="left", padx=6)

    # --- Format seçimi ---
    secenek_cerceve = ttk.LabelFrame(ana, text="Çıktı formatı", padding=8)
    secenek_cerceve.pack(fill="x", pady=4)

    format_degisken = tk.StringVar(value="word")
    ttk.Radiobutton(secenek_cerceve, text="Word (.docx)",
                    variable=format_degisken,
                    value="word").pack(side="left", padx=10)
    ttk.Radiobutton(secenek_cerceve, text="Excel (.xlsx)",
                    variable=format_degisken,
                    value="excel").pack(side="left", padx=10)

    ocr_var = tk.BooleanVar(value=True)
    ocr_kutusu = ttk.Checkbutton(
        secenek_cerceve,
        text="Taranmış sayfalar için OCR",
        variable=ocr_var)
    ocr_kutusu.pack(side="left", padx=10)
    if not ocr_kullanilabilir():
        ocr_var.set(False)
        ocr_kutusu.config(state="disabled",
                          text="OCR kullanılamıyor (tesseract yok)")

    # --- Çıktı klasörü ---
    cikti_cerceve = ttk.Frame(ana)
    cikti_cerceve.pack(fill="x", pady=4)
    cikti_degisken = tk.StringVar(value="")

    def klasor_sec():
        klasor = filedialog.askdirectory(title="Çıktı klasörünü seçin")
        if klasor:
            cikti_degisken.set(klasor)

    ttk.Button(cikti_cerceve, text="📁 Çıktı Klasörü",
               command=klasor_sec).pack(side="left")
    cikti_etiket = ttk.Label(cikti_cerceve,
                             textvariable=cikti_degisken,
                             foreground="gray")
    cikti_etiket.pack(side="left", padx=8)
    ttk.Label(cikti_cerceve,
              text="(boş bırakılırsa PDF'in yanına kaydedilir)",
              foreground="gray").pack(side="left")

    # --- İlerleme ve durum ---
    ilerleme_cubugu = ttk.Progressbar(ana, mode="determinate")
    ilerleme_cubugu.pack(fill="x", pady=(10, 4))
    durum_degisken = tk.StringVar(value="Hazır")
    ttk.Label(ana, textvariable=durum_degisken).pack()

    donusturme_aktif = [False]

    def calistir():
        if donusturme_aktif[0]:
            return
        if not secili_dosyalar:
            messagebox.showwarning("Uyarı", "Lütfen önce PDF dosyası ekleyin.")
            return

        donusturme_aktif[0] = True
        baslat_dugme.config(state="disabled")
        dosyalar = list(secili_dosyalar)
        secim = format_degisken.get()
        cikti = cikti_degisken.get() or None
        ocr_secimi = ocr_var.get()

        def arka_plan():
            hatalar = []
            basarili = []
            toplam_dosya = len(dosyalar)
            for i, pdf in enumerate(dosyalar):
                ad = os.path.basename(pdf)

                def sayfa_ilerleme(sayfa, toplam_sayfa, i=i, ad=ad):
                    yuzde = (i + sayfa / max(toplam_sayfa, 1)) / toplam_dosya * 100
                    pencere.after(0, lambda: (
                        ilerleme_cubugu.config(value=yuzde),
                        durum_degisken.set(
                            f"{ad} — sayfa {sayfa}/{toplam_sayfa}")))

                try:
                    hedef = donustur(pdf, secim, cikti, sayfa_ilerleme,
                                     ocr=ocr_secimi)
                    basarili.append(str(hedef))
                except Exception as hata:
                    hatalar.append(f"{ad}: {hata}")
                    traceback.print_exc()

            def bitir():
                donusturme_aktif[0] = False
                baslat_dugme.config(state="normal")
                ilerleme_cubugu.config(value=100 if not hatalar else 0)
                if hatalar:
                    durum_degisken.set("Bazı dosyalar dönüştürülemedi.")
                    messagebox.showerror(
                        "Hata",
                        "Şu dosyalar dönüştürülemedi:\n\n" + "\n".join(hatalar))
                else:
                    durum_degisken.set(
                        f"Tamamlandı! {len(basarili)} dosya dönüştürüldü.")
                    messagebox.showinfo(
                        "Başarılı",
                        "Dönüştürme tamamlandı:\n\n" + "\n".join(basarili))

            pencere.after(0, bitir)

        threading.Thread(target=arka_plan, daemon=True).start()

    baslat_dugme = ttk.Button(ana, text="🔄 Dönüştür", command=calistir)
    baslat_dugme.pack(pady=10, ipadx=20, ipady=4)

    pencere.mainloop()


# ---------------------------------------------------------------------------
# Komut satırı
# ---------------------------------------------------------------------------

def main():
    ayristirici = argparse.ArgumentParser(
        description="PDF dosyalarını Word veya Excel formatına çevirir.")
    ayristirici.add_argument("pdf", nargs="*",
                             help="Dönüştürülecek PDF dosyaları")
    ayristirici.add_argument("--format", choices=["word", "excel"],
                             default="word",
                             help="Çıktı formatı (varsayılan: word)")
    ayristirici.add_argument("--cikti", default=None,
                             help="Çıktı klasörü (varsayılan: PDF'in yanı)")
    ayristirici.add_argument("--ocr-kapali", action="store_true",
                             help="Taranmış sayfalar için OCR kullanma")
    argumanlar = ayristirici.parse_args()

    if not argumanlar.pdf:
        arayuz_baslat()
        return

    for pdf in argumanlar.pdf:
        print(f"Dönüştürülüyor: {pdf} → {argumanlar.format} ...")
        hedef = donustur(pdf, argumanlar.format, argumanlar.cikti,
                         ocr=not argumanlar.ocr_kapali)
        print(f"  Oluşturuldu: {hedef}")


if __name__ == "__main__":
    main()
